from __future__ import annotations

import argparse
import csv
import signal
from datetime import date as _date
from datetime import timedelta as _td
from pathlib import Path

from .browser import inspect_report
from .collector import daily_date, download_report_xlsx, weekly_date_range
from .exporter import export_csv, export_xlsx
from .filter_config import load_item_classification_filter
from .settings import config_dir, required_dir, required_env, today_kst
from .transform import filter_and_deduplicate


def _read_rows(path: Path) -> list[dict[str, str]]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise RuntimeError("XLSX input requires openpyxl; install project dependencies first") from error
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            all_rows = list(sheet.iter_rows(values_only=True))
            # G2B MicroStrategy exports have metadata rows before the actual data header.
            # Locate the header row by finding '입력일자' in the first column.
            header_idx = 0
            for i, row in enumerate(all_rows):
                if row and row[0] is not None and str(row[0]).strip() == "입력일자":
                    header_idx = i
                    break
            headers = [str(v).strip() if v is not None else "" for v in all_rows[header_idx]]
            return [
                dict(zip(headers, row, strict=False))
                for row in all_rows[header_idx + 1 :]
                if any(v is not None for v in row)
            ]
        finally:
            workbook.close()
    with path.open(encoding="utf-8-sig", newline="") as source:
        return list(csv.DictReader(source))


def _config_file(given: Path | None, name: str) -> Path:
    """--filter-file / --products-file 이 없으면 JODAL_CONFIG_DIR 의 같은 이름 파일 (설정은 기밀 — 저장소 밖에서 제자리 읽기)."""
    return given if given is not None else config_dir() / name


def _terminate(signum, _frame) -> None:
    """systemd 중지(타임아웃·stop·재부팅)의 SIGTERM 을 SystemExit 로 — job_runs 행이 'running' 으로 남지 않고 failed 로 닫힌다."""
    raise SystemExit(128 + signum)


def main(argv: list[str] | None = None) -> None:
    signal.signal(signal.SIGTERM, _terminate)
    parser = argparse.ArgumentParser(prog="jodal")
    commands = parser.add_subparsers(dest="command", required=True)
    inspect = commands.add_parser("inspect", help="open report once and save safe response metadata")
    inspect.add_argument("--output-dir", type=Path, default=Path("data/inspection"))
    inspect.add_argument("--headless", action="store_true")
    export = commands.add_parser("filter-export", help="filter a downloaded CSV/XLSX and export it")
    export.add_argument("--input", type=Path, required=True)
    export.add_argument("--filter-file", type=Path, default=None, help="기본: $JODAL_CONFIG_DIR/filter.md")
    export.add_argument("--output", type=Path, required=True)
    export.add_argument("--format", choices=("csv", "xlsx"), default="csv")
    collect = commands.add_parser("collect", help="download the public report and create a filtered XLSX")
    collect.add_argument("--output-dir", type=Path, default=Path("data/exports"))
    collect.add_argument("--filter-file", type=Path, default=None, help="기본: $JODAL_CONFIG_DIR/filter.md")
    collect.add_argument("--headed", action="store_true", help="show the browser while collecting")
    daily_cmd = commands.add_parser("daily", help="등록물품·납품요구를 수집해 Prami DB에 기록하고 보고 메일을 보낸다")
    daily_cmd.add_argument("--output-dir", type=Path, default=Path("data/exports"))
    daily_cmd.add_argument("--filter-file", type=Path, default=None, help="기본: $JODAL_CONFIG_DIR/filter.md")
    daily_cmd.add_argument("--headed", action="store_true")
    daily_cmd.add_argument("--date", help="기준 종료일 YYYY-MM-DD (기본: 어제, D-1). 이날부터 7일 이전까지 조회")
    daily_cmd.add_argument(
        "--products-file", type=Path, default=None, help="제품별 키워드 파일 (기본: $JODAL_CONFIG_DIR/products.md)"
    )
    daily_cmd.add_argument("--no-email", action="store_true", help="메일을 보내지 않는다 (시험용)")
    daily_cmd.add_argument(
        "--skip-if-done",
        action="store_true",
        help="이 기준일 보고를 이미 마쳤으면 아무것도 하지 않는다 (재시도 일정용)",
    )
    bids_cmd = commands.add_parser(
        "bids", help="나라장터 입찰공고 중 제품 키워드 관련 공고를 Prami DB(bid_notices)에 반영"
    )
    bids_cmd.add_argument("--days", type=int, default=14, help="오늘부터 며칠 전 공고까지 조회 (기본 14일)")
    bids_cmd.add_argument("--since", default=None, help="이 날짜(YYYY-MM-DD)부터 조회 — 처음 채울 때")
    args = parser.parse_args(argv)
    try:
        # 시작 시 필수 환경변수 확인 — 수집·DB 접속 전에 멈춘다
        if args.command in ("daily", "bids"):
            required_env("COLLECTOR_DATABASE_URL")
        if args.command == "daily":
            required_dir("JODAL_REF_DIR")
            config_dir()
        if args.command in ("daily", "collect", "filter-export"):
            args.filter_file = _config_file(args.filter_file, "filter.md")
        if args.command == "daily":
            args.products_file = _config_file(args.products_file, "products.md")
    except RuntimeError as error:
        parser.exit(2, f"jodal: 시작 오류: {error}\n")
    if args.command == "bids":
        from .db import load_bid_rules, sync_bids
        from .g2b_bid import fetch_bid_notices

        end = today_kst()
        days = args.days
        if args.since:
            days = (end - _date.fromisoformat(args.since)).days + 1
        rules, has_new = load_bid_rules()
        if has_new:  # 화면에서 키워드를 새로 넣었으면 올해 공고까지 거슬러 찾는다
            days = max(days, (end - _date(end.year, 1, 1)).days + 1)
            print(f"[입찰공고] 새 키워드가 있어 {end.year}-01-01부터 다시 찾습니다")
        notices = fetch_bid_notices(end - _td(days=days - 1), end, rules)
        added, updated = sync_bids(notices, rules)
        print(
            f"[입찰공고] {end - _td(days=days - 1)} ~ {end}: 관련 공고 {len(notices)}건 · 새로 추가 {added}건 · 갱신 {updated}건"
        )
        return
    if args.command == "inspect":
        result = inspect_report(output_dir=args.output_dir, headed=not args.headless)
        print(f"opened: {result.final_url}")
        print(f"responses recorded: {len(result.response_urls)}")
        return
    if args.command == "collect":
        start_date, end_date = weekly_date_range()
        stamp = f"{start_date:%Y%m%d}_to_{end_date:%Y%m%d}"
        download_path = args.output_dir / "downloads" / f"g2b_items_{stamp}.xlsx"
        downloaded = download_report_xlsx(
            destination=download_path, start_date=start_date, end_date=end_date, headed=args.headed
        )
        item_filter = load_item_classification_filter(args.filter_file)
        records = filter_and_deduplicate(_read_rows(downloaded), item_filter.code, item_filter.item_keywords)
        output = args.output_dir / f"g2b_items_filtered_{stamp}.xlsx"
        count = export_xlsx(records, output)
        print(f"downloaded: {downloaded}")
        print(f"exported {count} filtered rows: {output}")
        return
    if args.command == "daily":
        from .db import run_daily

        end_date = _date.fromisoformat(args.date) if args.date else daily_date()
        run_daily(
            end_date=end_date,
            headed=args.headed,
            no_email=args.no_email,
            skip_if_done=args.skip_if_done,
            filter_file=args.filter_file,
            products_file=args.products_file,
            read_rows=_read_rows,
            output_dir=args.output_dir,
        )
        return
    item_filter = load_item_classification_filter(args.filter_file)
    records = filter_and_deduplicate(_read_rows(args.input), item_filter.code, item_filter.item_keywords)
    count = export_xlsx(records, args.output) if args.format == "xlsx" else export_csv(records, args.output)
    print(
        f"exported {count} rows for 물품분류번호={item_filter.code}, keywords={','.join(item_filter.item_keywords)}: {args.output}"
    )


if __name__ == "__main__":
    main()
